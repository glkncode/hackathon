"""Real public-API Docker E2E. Creates labelled synthetic fixtures; never deletes volumes.

Run from repository root: ai-service/.venv/Scripts/python.exe scripts/template_e2e.py
Requires running Compose and admin credentials in .env (never printed).
--restart exercises service restart and Python outage; it briefly interrupts local services.
"""
import argparse
import hashlib
import io
import json
import secrets
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from zipfile import ZipFile

import httpx
from docx import Document
from dotenv import dotenv_values
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
BASE = "http://localhost:8080"
MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def request(client, method, path, expected=200, csrf=True, **kwargs):
    headers = {"Origin": BASE}
    if csrf and client.cookies.get("ai_csrf"):
        headers["X-CSRF-Token"] = client.cookies.get("ai_csrf")
    response = client.request(method, path, headers=headers, **kwargs)
    assert response.status_code == expected, f"{method} {path}: expected {expected}, got {response.status_code} {response.text[:300]}"
    return response


def session(username, password):
    client = httpx.Client(base_url=BASE, timeout=90, trust_env=False)
    request(client, "POST", "/api/v1/auth/login", json={"username": username, "password": password})
    return client


def compose(*args):
    subprocess.run(["docker", "compose", *args], cwd=ROOT, check=True, capture_output=True, timeout=180)


def health():
    for _ in range(60):
        try:
            if httpx.get(BASE + "/healthz", timeout=2, trust_env=False).status_code == 200:
                return
        except httpx.HTTPError:
            pass
        time.sleep(1)
    raise AssertionError("API did not become healthy")


def fill(client, draft, values):
    path = "/api/v1/template-drafts/" + draft["id"]
    while True:
        validation = request(client, "GET", path + "/next-question").json()
        if validation["ready"]:
            return draft
        field = validation["next_question"]["key"]
        draft = request(client, "POST", path + "/answers", json={
            "field_key": field, "value": values[field], "expected_revision": draft["revision"]}).json()


def verify_docx(original, output, values):
    with ZipFile(io.BytesIO(original)) as source, ZipFile(io.BytesIO(output)) as result:
        assert result.testzip() is None
        assert set(source.namelist()) == set(result.namelist())
        for name in result.namelist():
            if name.endswith(".xml"):
                etree.fromstring(result.read(name))
            if name != "word/document.xml":
                assert source.read(name) == result.read(name), name
        a = etree.fromstring(source.read("word/document.xml"))
        b = etree.fromstring(result.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        for selector in ["//w:pPr", "//w:rPr", "//w:tblPr", "//w:tcPr", "//w:sectPr"]:
            assert [etree.tostring(n) for n in a.xpath(selector, namespaces=ns)] == [
                etree.tostring(n) for n in b.xpath(selector, namespaces=ns)], selector
    d = Document(io.BytesIO(output))
    assert len(d.tables) == 4
    assert values["group_subject"] in d.paragraphs[5].text
    assert values["group_subject"] in d.paragraphs[8].text
    assert values["purpose"] in d.paragraphs[6].text
    assert d.tables[2].cell(1, 1).text == values["leader"]
    assert d.tables[2].cell(2, 1).text == values["secretary"]
    assert d.tables[2].cell(3, 1).text == "\n".join(values["members"])
    assert "15.12.2026" in d.paragraphs[9].text
    assert values["actions"] in d.paragraphs[9].text
    assert values["result"] in d.paragraphs[9].text
    assert values["termination"] in d.paragraphs[12].text
    assert values["controller"] in d.paragraphs[13].text
    assert d.tables[3].cell(0, 2).text == values["signer"]
    assert "НЕ УДАЛЯТЬ" in d.paragraphs[3].text
    assert "ВИЗУАЛИЗАЦИЯ ЭП (НЕ УДАЛЯТЬ)" in d.tables[3].cell(1, 0).text
    assert "____" not in "\n".join(p.text for p in d.paragraphs[5:14])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--restart", action="store_true")
    args = parser.parse_args()
    env = dotenv_values(ROOT / ".env")
    health()
    admin = session(env.get("ADMIN_USERNAME", "admin"), env["ADMIN_PASSWORD"])
    suffix = secrets.token_hex(4)
    users = []
    for index in range(2):
        username, password = f"template_e2e_{suffix}_{index}", "Synthetic-" + secrets.token_hex(16)
        user = request(admin, "POST", "/api/v1/users", 201, json={"username": username, "password": password}).json()
        users.append((user, session(username, password)))
    (user, client), (_, foreign) = users
    source_path = next((ROOT / "data/demo").glob("*рабоч*группы.docx"))
    original = source_path.read_bytes()
    doc = request(admin, "POST", "/api/v1/documents", 201, data={"title": "E2E working group " + suffix},
                  files={"file": (source_path.name, original, MIME)}).json()
    fields = json.loads((ROOT / "docs/template-profiles/working-group.json").read_text(encoding="utf-8-sig"))
    template = request(admin, "POST", "/api/v1/templates", 201, json={
        "document_id": doc["id"], "document_version_id": doc["latest_version"]["id"],
        "name": "E2E Распоряжение " + suffix, "description": "Синтетическая проверка", "fields": fields}).json()
    access = f'/api/v1/documents/{doc["id"]}/access'
    request(admin, "POST", access, 201, json={"user_id": user["id"]})
    assert template["id"] in [x["id"] for x in request(client, "GET", "/api/v1/templates?limit=100").json()["items"]]
    create = f'/api/v1/templates/{template["id"]}/drafts'
    request(foreign, "POST", create, 404)
    request(client, "POST", create, 403, csrf=False)
    with httpx.Client(base_url=BASE, trust_env=False) as anonymous:
        request(anonymous, "GET", "/api/v1/templates", 401)
    draft = request(client, "POST", create, 201).json()
    path = f'/api/v1/template-drafts/{draft["id"]}'
    request(foreign, "GET", path, 404)
    request(admin, "GET", path, 404)
    request(client, "GET", "/api/v1/template-drafts/00000000-0000-0000-0000-000000000001", 404)
    request(client, "POST", path + "/generate", 422, json={"expected_revision": 1})
    values = dict(approval_person="Уполномоченное тестовое лицо", approval_date="2026-10-24",
                  organization="Тестовая организация", group_subject="Внедрение электронного архива",
                  purpose="Подготовка предложений по внедрению электронного архива документов",
                  leader="Иванов И.И., начальник отдела цифровизации",
                  secretary="Петрова А.А., специалист отдела ДОУ",
                  members=["Сидоров С.С., специалист ИТ", "Кузнецова Е.В., специалист архивного подразделения"],
                  deadline="2026-12-15", actions="Проанализировать текущие процессы хранения документов и подготовить план внедрения",
                  result="План мероприятий по внедрению электронного архива на рассмотрение руководства",
                  termination="представления плана мероприятий", controller="Смирнов А.В., начальник управления",
                  signer="Фамилия, инициалы уполномоченного подписанта, заданные в тесте")
    first = fields[0]["key"]
    draft = request(client, "POST", path + "/answers", json={"expected_revision": 1, "field_key": first, "value": values[first]}).json()
    request(client, "PATCH", path, 409, json={"expected_revision": 1, "values": {first: "stale"}})
    request(client, "PATCH", path, 400, json={"expected_revision": draft["revision"], "values": {"unknown": "x"}})
    assert len(request(client, "POST", path + "/validate").json()["missing_fields"]) == len(fields) - 1
    print("PASS: auth, CSRF, ownership, incomplete draft, revision conflict, unknown keys", flush=True)
    if args.restart:
        compose("restart", "api", "ai-service")
        health()
        assert request(client, "GET", path).json()["values"][first] == values[first]
    draft = fill(client, draft, values)
    assert draft["status"] == "ready_to_generate"
    assert request(client, "POST", path + "/validate").json()["ready"]
    def generate(_):
        return request(client, "POST", path + "/generate", json={"expected_revision": draft["revision"]}).json()
    with ThreadPoolExecutor(max_workers=2) as pool:
        generated = list(pool.map(generate, range(2)))
    assert generated[0]["generated_document_id"] == generated[1]["generated_document_id"]
    generated = generated[0]
    result_id = generated["generated_document_id"]
    download = generated["download_url"]
    response = request(client, "GET", download)
    assert response.headers["content-type"] == MIME
    output = response.content
    verify_docx(original, output, values)
    metadata = request(client, "GET", f"/api/v1/generated-documents/{result_id}").json()
    assert "storage_key" not in metadata and "storage_path" not in metadata
    assert metadata["sha256"] == hashlib.sha256(output).hexdigest()
    request(foreign, "GET", download, 404)
    request(foreign, "GET", f"/api/v1/generated-documents/{result_id}", 404)
    print("PASS: wizard, parallel idempotency, real DOCX values/tables/styles/ZIP/service elements", flush=True)
    request(admin, "DELETE", access + "/" + user["id"], 204)
    for method, endpoint, kwargs in [("GET", path, {}), ("GET", download, {}),
                                     ("POST", path + "/generate", {"json": {"expected_revision": draft["revision"]}})]:
        request(client, method, endpoint, 404, **kwargs)
    assert draft["id"] not in [x["id"] for x in request(client, "GET", "/api/v1/template-drafts?limit=100").json()["items"]]
    request(admin, "POST", access, 201, json={"user_id": user["id"]})
    assert request(client, "GET", download).content == output
    print("PASS: live ACL revoke/restore, direct UUID download denied", flush=True)
    if args.restart:
        compose("restart", "api", "ai-service")
        health()
        assert request(client, "GET", download).content == output
        # A new revision forces rendering, rather than exercising an idempotent cache hit.
        draft = request(client, "PATCH", path, json={"expected_revision": draft["revision"], "values": {"purpose": values["purpose"]}}).json()
        compose("stop", "ai-service")
        try:
            request(client, "POST", path + "/generate", 502, json={"expected_revision": draft["revision"]})
            assert request(client, "GET", path).json()["values"] == values
            assert request(client, "GET", download).content == output
        finally:
            compose("start", "ai-service")
        for _ in range(45):
            check = client.post(path + "/generate", json={"expected_revision": draft["revision"]},
                                headers={"Origin": BASE, "X-CSRF-Token": client.cookies.get("ai_csrf")})
            if check.status_code == 200:
                break
            assert check.status_code == 502
            time.sleep(1)
        else:
            raise AssertionError("Python did not recover")
        assert check.json()["generated_document_id"] != result_id
        print("PASS: partial draft restart, result restart, Python outage/retry, independent revision result", flush=True)
    out = ROOT / ".local-integration/template-e2e"
    out.mkdir(parents=True, exist_ok=True)
    (out / "working-group-filled.docx").write_bytes(output)
    report = dict(template_id=template["id"], draft_id=draft["id"], generated_document_id=result_id,
                  sha256=hashlib.sha256(output).hexdigest(), restart_and_outage=args.restart,
                  source_sha256=hashlib.sha256(original).hexdigest(), fields=len(fields))
    (out / "report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    assert source_path.read_bytes() == original
    print("PASS: source unchanged; saved .local-integration/template-e2e/working-group-filled.docx", flush=True)


if __name__ == "__main__":
    main()
