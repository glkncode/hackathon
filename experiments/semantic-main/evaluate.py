"""Explicit main-local verification, only localhost:8080 and main containers.

Artifacts with real document/user contents stay in ignored .local-integration.
Creates one named synthetic member and seven documents; never deletes old data.
"""
import sys, json, time, subprocess, hashlib, threading
from pathlib import Path
from uuid import uuid4
import httpx
import numpy as np

HERE=Path(__file__).resolve().parent;ROOT=HERE.parent.parent
sys.path.insert(0,str(ROOT/'experiments/rag'))
import benchmark as fixtures
OUT=ROOT/'.local-integration/main-semantic-20261024/e2e';OUT.mkdir(exist_ok=True)
BASE='http://localhost:8080'
PASSWORD='main-local-fixture-only-20261024'
def save(name,data): (OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
def load(name): return json.loads((OUT/name).read_text(encoding='utf-8'))
def sql(s):
    p=subprocess.run(['docker','exec','-i','hackathon-postgres-1','psql','-X','-v','ON_ERROR_STOP=1','-U','app','-d','ai_assistant','-At'],input=s,text=True,encoding='utf-8',capture_output=True,check=True)
    return json.loads(p.stdout)
def admin_password():
    d=json.loads(subprocess.check_output(['docker','inspect','hackathon-api-1']))[0]
    return next(e.split('=',1)[1] for e in d['Config']['Env'] if e.startswith('ADMIN_PASSWORD='))
class API:
    def __init__(self,admin=False):
        self.http=httpx.Client(base_url=BASE,headers={'Origin':BASE},trust_env=False,timeout=120)
        self.user=self.request('POST','/api/v1/auth/login',json={'username':'admin' if admin else 'semantic-main-check','password':admin_password() if admin else PASSWORD})['user']
        self.http.headers['X-CSRF-Token']=self.http.cookies['ai_csrf']
    def request(self,method,path,**kw):
        r=self.http.request(method,path,**kw);r.raise_for_status();return r.json() if r.content else None
    def search(self,q):return self.request('POST','/api/v1/search',json={'query':q,'limit':10})['results']
    def turn(self,q,session=None,key=None):
        if not session:session=self.request('POST','/api/v1/chat/sessions',json={'title':'Main local verification'})['id']
        key=key or str(uuid4());turn=self.request('POST','/api/v1/chat/sessions/'+session+'/messages',json={'content':q,'client_message_id':key})
        return session,key,turn
    def close(self):self.request('POST','/api/v1/auth/logout');self.http.close()
def wait_index():
    for _ in range(180):
        state=sql("SELECT json_build_object('total',count(*),'ready',count(*) FILTER(WHERE s.status='ready')) FROM documents d JOIN document_versions v ON v.id=d.active_version_id LEFT JOIN semantic_indexes s ON s.document_version_id=v.id WHERE v.status='ready'")
        if state['total']==state['ready']:return state
        time.sleep(1)
    raise RuntimeError('Index readiness timeout')
fixtures.OUT=OUT;fixtures.sql=sql
def prepare():
    assert not (OUT/'manifest.json').exists()
    corpus=json.loads((HERE/'corpus.json').read_text(encoding='utf-8'))
    regression=json.loads((ROOT/'experiments/semantic-pilot/corpus-final.json').read_text(encoding='utf-8'))['documents'][0]
    corpus['documents'].append(regression)
    save('corpus.json',corpus)
    admin=API(True)
    if (OUT/'prepare-state.json').exists(): state=load('prepare-state.json')
    else: state={'member':admin.request('POST','/api/v1/users',json={'username':'semantic-main-check','password':PASSWORD}),'documents':{}};save('prepare-state.json',state)
    for d in corpus['documents']:
        if d['key'] in state['documents']:continue
        path=fixtures.document_file(d)
        with path.open('rb') as f: doc=admin.request('POST','/api/v1/documents',data={'title':d['title']},files={'file':(path.name,f)})
        fixtures.wait_ready(doc['latest_version']['id'])
        admin.request('POST',f"/api/v1/documents/{doc['id']}/access",json={'user_id':state['member']['id']})
        state['documents'][d['key']]={**doc,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()};save('prepare-state.json',state)
        print('main ingestion ready',d['key'],flush=True)
    original=sql("SELECT row_to_json(d) FROM (SELECT id,title,active_version_id FROM documents WHERE id='9a7d2e68-739f-4cf9-a935-c952e32db73b') d")
    if not sql(f"SELECT count(*) FROM document_access WHERE document_id='{original['id']}' AND user_id='{state['member']['id']}'"):
        admin.request('POST',f"/api/v1/documents/{original['id']}/access",json={'user_id':state['member']['id']})
    chunks=fixtures.visible_chunks(state['member']['id'])
    questions=[]
    for d in corpus['documents']:
        for q in d['questions']:
            if d['key']=='meeting_rooms' and q['id']!='f03':continue
            questions.append({**q,'key':d['key'],'split':'regression' if d['key']=='meeting_rooms' else 'independent'})
    questions += [{**q,'split':'negative'} for q in corpus['negatives']]
    questions += [
      {'id':'old01','split':'original','query':'Кто должен распределять задачи рабочей группы и следить за сроками?','gold':['Руководителю Рабочей группы'],'facts':['руководител','распредел','срок']},
      {'id':'old02','split':'original','query':'Кто оформляет итоги заседаний рабочей группы и хранит рабочие материалы?','gold':['Секретарю Рабочей группы'],'facts':['секретар','оформл','хран']},
    ]
    for q in questions:
        q['gold_chunk_ids']=[c['chunk_id'] for c in chunks if (not q.get('key') or c['document_id']==state['documents'][q['key']]['id']) and any(' '.join(g.split()) in ' '.join(c['content'].split()) for g in q['gold'])]
        assert not q['gold'] or q['gold_chunk_ids'],q['id']
    save('manifest.json',{'member_id':state['member']['id'],'documents':state['documents'],'original':original,'chunks':chunks,'questions':questions,'accessible_documents':len({c['document_id'] for c in chunks}),'frozen_before_search':True})
    assert len({c['document_id'] for c in chunks})==8
    save('indexing.json',wait_index());admin.close()
    print('Main manifest frozen: eight accessible documents.',flush=True)
class Resources:
    def __enter__(self):
        self.stop=threading.Event();self.samples=[]
        def sample():
            while not self.stop.is_set():
                p=subprocess.run(['docker','stats','--no-stream','--format','{{json .}}','hackathon-api-1','hackathon-ai-service-1','hackathon-postgres-1'],capture_output=True,text=True)
                self.samples.extend(json.loads(l) for l in p.stdout.splitlines());self.stop.wait(3)
        self.thread=threading.Thread(target=sample,daemon=True);self.thread.start();return self
    def __exit__(self,*args):self.stop.set();self.thread.join();save('resources-'+sys.argv[1]+'.json',self.samples)
def evaluate(mode):
    info=json.loads(subprocess.check_output(['docker','inspect','hackathon-api-1']))[0]
    assert 'SEMANTIC_SEARCH_ENABLED='+('true' if mode=='semantic' else 'false') in info['Config']['Env']
    manifest=load('manifest.json');allowed={c['chunk_id'] for c in manifest['chunks']}
    name='observations-'+mode+'.json';rows=load(name) if (OUT/name).exists() else []
    done={r['id'] for r in rows};api=API()
    with Resources():
        for q in manifest['questions']:
            if q['id'] in done:continue
            start=time.perf_counter();sources=api.search(q['query']);search_ms=(time.perf_counter()-start)*1000
            start=time.perf_counter();session,key,turn=api.turn(q['query']);chat_ms=(time.perf_counter()-start)*1000
            assert {s['chunk_id'] for s in sources+turn['sources']}<=allowed
            assert len({s['chunk_id'] for s in turn['sources']})==len(turn['sources'])
            rows.append({'id':q['id'],'split':q['split'],'query':q['query'],'gold':q['gold_chunk_ids'],'facts':q.get('facts',[]),'search':sources,'search_ms':search_ms,'chat':turn,'chat_ms':chat_ms,'session_id':session})
            save(name,rows);print(mode,q['id'],len(sources),turn['mode'],flush=True)
    api.close()
def activate():
    assert (OUT.parent/'main-before-activation.json').exists()
    assert wait_index()['total']==8
    path=ROOT/'.env';text=path.read_text(encoding='utf-8');assert 'SEMANTIC_SEARCH_ENABLED=false' in text
    path.write_text(text.replace('SEMANTIC_SEARCH_ENABLED=false','SEMANTIC_SEARCH_ENABLED=true'),encoding='utf-8')
    subprocess.run(['docker','compose','up','-d','--no-deps','--no-build','api'],cwd=ROOT,check=True)
    print('Main search enabled after verified restore and complete backfill.')
def security():
    manifest=load('manifest.json');api=API();admin=API(True);mid=manifest['member_id'];doc=manifest['documents']['keys']
    evidence={}
    # Every citation points to the same original stored file via normal Go API.
    downloads=[]
    for d in list(manifest['documents'].values())+[manifest['original']]:
        version=d.get('active_version_id') or d['latest_version']['id']
        r=api.http.get(f"/api/v1/documents/{d['id']}/versions/{version}/download");r.raise_for_status()
        actual=hashlib.sha256(r.content).hexdigest()
        expected=d.get('sha256') or (OUT.parent/'document-files.sha256').read_text().split()[0]
        assert actual==expected
        downloads.append({'document_id':d['id'],'version':version,'sha256':actual})
    evidence['source_downloads']=downloads
    # A real contextual follow-up, recorded without changing understanding rules.
    session,key,first=api.turn('К кому обращаться за запасным ключом от кабинета?')
    _,_,follow=api.turn('А до какого времени его нужно вернуть?',session)
    evidence['clarification']={'session':session,'first':first,'follow_up':follow}
    # Retry the exact public Chat request, which must not add messages.
    session,key,first=api.turn('Кто отвечает за полив демонстрационных растений?')
    path='/api/v1/chat/sessions/'+session+'/messages'
    retry=api.request('POST',path,json={'content':'Кто отвечает за полив демонстрационных растений?','client_message_id':key})
    assert retry==first
    assert len(api.request('GET','/api/v1/chat/sessions/'+session)['messages'])==2
    evidence['chat_idempotency']=True
    row=next(r for r in load('observations-semantic.json') if r['id']=='m01')
    admin.request('DELETE',f"/api/v1/documents/{doc['id']}/access/{mid}")
    try:
        assert not any(s['document_id']==doc['id'] for s in api.search(row['query']))
        history=api.request('GET','/api/v1/chat/sessions/'+row['session_id'])
        assert any(m['content']=='Содержимое ответа недоступно из-за изменения прав доступа к документам.' for m in history['messages'] if m['role']=='assistant')
        r=api.http.get(f"/api/v1/documents/{doc['id']}/versions/{doc['latest_version']['id']}/download")
        assert r.status_code in (403,404)
        evidence['acl_revoke']={'search_hidden':True,'history_redacted':True,'download_status':r.status_code}
    finally:admin.request('POST',f"/api/v1/documents/{doc['id']}/access",json={'user_id':mid})
    # Existing authentication and CSRF remain enforced.
    anon=httpx.get(BASE+'/api/v1/documents',trust_env=False);assert anon.status_code==401
    bad=api.http.post('/api/v1/search',json={'query':'ключ','limit':10},headers={'X-CSRF-Token':'invalid'});assert bad.status_code==403
    evidence['auth_csrf']={'anonymous':anon.status_code,'bad_csrf':bad.status_code}
    # Version update affects only a synthetic document, after frozen comparison.
    from docx import Document
    path=OUT/'documents/keys-v2.docx';document=Document()
    document.add_paragraph('1. Запасной ключ от кабинета выдаёт дежурный администратор после проверки служебного удостоверения. Получатель расписывается в журнале.')
    document.add_paragraph('2. Ключ возвращают администратору не позднее 21:05 того же дня. Передавать ключ другому сотруднику запрещено.')
    document.save(path)
    with path.open('rb') as f:v=admin.request('POST',f"/api/v1/documents/{doc['id']}/versions",files={'file':(path.name,f)})
    fixtures.wait_ready(v['id']);wait_index()
    results=api.search('До какого времени возвращают запасной ключ от кабинета?')
    assert any(s['document_version_id']==v['id'] and '21:05' in s['content'] for s in results)
    assert not any(s['document_version_id']==doc['latest_version']['id'] for s in results)
    evidence['new_version']={'old':doc['latest_version']['id'],'new':v['id'],'search':results,'old_chunks_preserved':sql(f"SELECT count(*) FROM document_chunks WHERE document_version_id='{doc['latest_version']['id']}'")}
    save('security.json',evidence)
    # Python outage exercises real main Go fallback for both Search and Chat.
    subprocess.run(['docker','stop','hackathon-ai-service-1'],capture_output=True,check=True)
    try:
        fallback=api.search('плесени');assert fallback
        _,_,turn=api.turn('плесени');assert turn['mode']=='search' and turn['sources']
        evidence['python_outage']={'search':fallback,'chat':turn}
    finally:subprocess.run(['docker','start','hackathon-ai-service-1'],capture_output=True,check=True)
    save('security.json',evidence)
    api.close();admin.close()
    print('Main ACL/history/download, Auth/CSRF, idempotency, v2 and Python-outage FTS/Chat checks passed.')
def persistence():
    statement="SELECT json_build_object('count',count(*),'md5',md5(string_agg(row_data,E'\\n' ORDER BY row_data))) FROM (SELECT row_to_json(w)::text row_data FROM embedding_windows w) a"
    before=sql(statement)
    subprocess.run(['docker','compose','restart','postgres','ai-service','api'],cwd=ROOT,capture_output=True,check=True)
    for _ in range(90):
        try:
            after=sql(statement)
            if before==after:break
        except Exception:pass
        time.sleep(.5)
    else:raise RuntimeError('Vector persistence failed')
    assert wait_index()['total']==8
    save('persistence.json',{'before':before,'after':after,'identical_all_vector_rows':True,'indexing':wait_index()})
    print('All main vector row fingerprints identical after PostgreSQL/Python/Go restart.')
def summarize():
    report={}
    for mode in ('fts','semantic'):
        rows=load('observations-'+mode+'.json');positive=[r for r in rows if r['split']=='independent'];negative=[r for r in rows if r['split']=='negative']
        report[mode]={'positive_count':len(positive),'found_gold_search':sum(bool(set(r['gold']) & {s['chunk_id'] for s in r['search']}) for r in positive),'generated_with_gold':sum(r['chat']['mode']=='generated' and bool(set(r['gold']) & {s['chunk_id'] for s in r['chat']['sources']}) for r in positive),'negative_generated':sum(r['chat']['mode']=='generated' for r in negative),'negative_count':len(negative),'search_ms':dict(zip(['p50','p95'],np.percentile([r['search_ms'] for r in rows],[50,95]).tolist())),'chat_ms':dict(zip(['p50','p95'],np.percentile([r['chat_ms'] for r in rows],[50,95]).tolist()))}
    save('metrics.json',report);print(json.dumps(report,indent=2))
def regression():
    manifest=load('manifest.json');q=next(q for q in manifest['questions'] if q['id']=='f03')
    code="""import json,os,urllib.request,sys
body={'kind':'query','inputs':[sys.argv[1]]}
req=urllib.request.Request('http://127.0.0.1:8000/internal/v1/embeddings',data=json.dumps(body).encode(),headers={'Content-Type':'application/json','X-Service-Token':os.environ['AI_SERVICE_TOKEN']})
with urllib.request.urlopen(req) as r:print(json.dumps(json.load(r)['embeddings'][0]))
"""
    v=json.loads(subprocess.check_output(['docker','exec','hackathon-ai-service-1','python','-c',code,q['query']]))
    vector=json.dumps(v,separators=(',',':'));mid=manifest['member_id']
    raw=sql(f"""SELECT json_agg(x) FROM (
      SELECT c.id AS chunk_id,1-min(w.embedding OPERATOR(public.<=>) '{vector}'::public.vector) AS similarity
      FROM document_chunks c JOIN document_versions dv ON dv.id=c.document_version_id
      JOIN documents d ON d.active_version_id=dv.id JOIN embedding_windows w ON w.chunk_id=c.id
      JOIN semantic_indexes s ON s.id=w.index_id AND s.status='ready'
      JOIN users u ON u.id='{mid}' AND u.is_active AND u.role='member'
      WHERE dv.status='ready' AND EXISTS(SELECT 1 FROM document_access a WHERE a.document_id=d.id AND a.user_id=u.id)
      GROUP BY c.id ORDER BY similarity DESC,c.id LIMIT 10) x""")
    assert set(q['gold_chunk_ids']) & {r['chunk_id'] for r in raw}
    api=API();results=api.search(q['query']);assert results==[]
    control=api.search('Кто оформляет итоги заседаний рабочей группы и хранит рабочие материалы?')
    assert any(s['document_id']==manifest['original']['id'] for s in control)
    api.close()
    save('known-false-negative.json',{'query':q['query'],'gold':q['gold_chunk_ids'],'raw_acl_retrieval':raw,'actual_go_gated_search':results,'expected_limitation':True,'gate_unchanged':True,'post_restart_original_doc_search':True})
    print('Known f03 false negative reproduced on main: raw retrieval finds gold, gate abstains. Post-restart original DOCX retrieval passed.')
if __name__=='__main__':
    stage=sys.argv[1]
    if stage in ('fts','semantic'):evaluate(stage)
    else:{'prepare':prepare,'activate':activate,'security':security,'persistence':persistence,'summary':summarize,'regression':regression}[stage]()
