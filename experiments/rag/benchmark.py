"""Local, opt-in retrieval experiment against disposable Docker infrastructure.

No imports from the running application, no production config changes. Commands:
prepare -> baseline -> index -> compare -> answers -> verify. See README.md.
"""

import argparse
import concurrent.futures
import hashlib
import json
import os
import platform
import importlib.metadata
import shutil
from pathlib import Path
import subprocess
import threading
import time
from uuid import UUID, uuid4

import httpx
import numpy as np
import psutil

from embedding import DIMENSION, MODEL, REVISION, E5, cosine_candidates, metrics, rrf

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
BASE = "http://127.0.0.1:18081"
PYTHON = "http://127.0.0.1:18083"
PG = "hackathon-hybrid-test-pg"
OUT = ROOT / ".local-integration/hybrid"
MODEL_PATH = ROOT / ".local-integration/e5-small"
PASSWORD = "experiment-test-only-123"  # Disposable fixture only.


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def file_digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def save(name, value):
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def load(name):
    return json.loads((OUT / name).read_text(encoding="utf-8"))


def sql(statement):
    # Fixed disposable database, never DATABASE_URL or production fallback.
    result = subprocess.run(
        ["docker", "exec", "-i", PG, "psql", "-X", "-v", "ON_ERROR_STOP=1",
         "-U", "search", "-d", "search_test", "-At"],
        input=statement, encoding="utf-8", capture_output=True, check=True,
    )
    return json.loads(result.stdout.strip())


class API:
    def __init__(self, username="admin"):
        self.http = httpx.Client(base_url=BASE, timeout=90, trust_env=False,
                                 headers={"Origin": BASE})
        self.user = self.request("POST", "/api/v1/auth/login",
                                 json={"username": username, "password": PASSWORD})["user"]
        self.http.headers["X-CSRF-Token"] = self.http.cookies["ai_csrf"]

    def request(self, method, path, **kwargs):
        response = self.http.request(method, path, **kwargs)
        response.raise_for_status()
        return response.json() if response.content else None

    def search(self, query):
        return self.request("POST", "/api/v1/search", json={"query": query, "limit": 10})["results"]

    def close(self):
        try:
            self.request("POST", "/api/v1/auth/logout")
        finally:
            self.http.close()


def text(block):
    if isinstance(block, dict):
        return block["repeat"] * block["count"] + block["tail"]
    return block


def document_file(doc):
    files = OUT / "documents"
    files.mkdir(parents=True, exist_ok=True)
    path = files / (doc["key"] + "." + doc["format"])
    if doc["format"] == "docx":
        from docx import Document
        document = Document()
        for block in doc["blocks"]:
            if isinstance(block, list):
                table = document.add_table(rows=0, cols=len(block[0]))
                for row in block:
                    for cell, value in zip(table.add_row().cells, row, strict=True):
                        cell.text = value
            else:
                document.add_paragraph(text(block))
        document.save(path)
    elif doc["format"] == "pdf":
        import pymupdf
        document = pymupdf.open()
        for block in doc["blocks"]:
            page = document.new_page()
            remaining = page.insert_textbox(pymupdf.Rect(50, 50, 540, 780), text(block),
                                           fontfile="C:/Windows/Fonts/arial.ttf", fontname="corpus", fontsize=12)
            if remaining < 0:
                raise ValueError("Synthetic PDF text did not fit")
        document.save(path)
        document.close()
    else:
        path.write_text("\n\n".join(text(b) for b in doc["blocks"]), encoding="utf-8")
    return path


def wait_ready(version_id):
    version_id = str(UUID(version_id))
    for _ in range(120):
        row = sql(f"SELECT row_to_json(v) FROM (SELECT status FROM document_versions WHERE id='{version_id}') v")
        if row["status"] == "ready":
            return
        if row["status"] == "failed":
            raise RuntimeError("Synthetic document ingestion failed")
        time.sleep(0.25)
    raise RuntimeError("Ingestion deadline")


def visible_chunks(user_id):
    user_id = str(UUID(user_id))
    return sql(f"""
        SELECT coalesce(json_agg(x ORDER BY document_id,chunk_index),'[]') FROM (
          SELECT c.id AS chunk_id,c.document_version_id,d.id AS document_id,
                 d.title AS document_title,v.version_number,c.chunk_index,c.content,
                 c.page_number,c.section_number,c.section_title
          FROM document_chunks c JOIN document_versions v ON v.id=c.document_version_id
          JOIN documents d ON d.id=v.document_id AND d.active_version_id=v.id
          JOIN users u ON u.id='{user_id}' AND u.is_active AND u.role='member'
          WHERE v.status='ready' AND EXISTS (
            SELECT 1 FROM document_access a WHERE a.document_id=d.id AND a.user_id=u.id)
        ) x
    """)


def prepare():
    if (OUT / "manifest.json").exists():
        raise RuntimeError("Manifest already frozen; use a new disposable database/output for a new corpus")
    corpus = json.loads((HERE / "corpus.json").read_text(encoding="utf-8"))
    admin = API()
    state = load("prepare-state.json") if (OUT / "prepare-state.json").exists() else None
    member = state["member"] if state else admin.request("POST", "/api/v1/users", json={"username": "rag-experiment", "password": PASSWORD})
    documents = state["documents"] if state else {}
    for doc in corpus["documents"]:
        if doc["key"] in documents:
            continue
        path = document_file(doc)
        # Create a genuine historical ready version before publishing travel v2.
        old = None
        if doc["key"] == "travel":
            old = admin.request("POST", "/api/v1/documents", data={"title": doc["title"]},
                                files={"file": ("old.txt", "1. Архивный тариф поездки — 17 рублей. Код архивного тарифа — 424242.".encode(), "text/plain")})
            wait_ready(old["latest_version"]["id"])
        with path.open("rb") as source:
            if old:
                version = admin.request("POST", f'/api/v1/documents/{old["id"]}/versions',
                                        files={"file": (path.name, source)})
                record = {"id": old["id"], "latest_version": version,
                          "historical_version_id": old["latest_version"]["id"]}
            else:
                record = admin.request("POST", "/api/v1/documents", data={"title": doc["title"]},
                                       files={"file": (path.name, source)})
        wait_ready(record["latest_version"]["id"])
        record["key"] = doc["key"]
        record["sha256"] = file_digest(path)
        documents[doc["key"]] = record
        if not doc.get("denied"):
            admin.request("POST", f'/api/v1/documents/{record["id"]}/access', json={"user_id": member["id"]})
        save("prepare-state.json", {"member": member, "documents": documents})
        print("ingested", doc["key"], record["latest_version"]["id"], flush=True)
    chunks = visible_chunks(member["id"])
    questions = []
    for doc in corpus["documents"]:
        for question in doc["questions"]:
            questions.append({**question, "document_key": doc["key"], "split": doc["split"]})
    questions += corpus["negatives"]
    questions.append({"id": "hn07", "split": "holdout", "query": "Какой код архивного тарифа поездки?",
                      "gold": [], "kind": "historical_version"})
    for q in questions:
        ids = []
        for needle in q["gold"]:
            key = q.get("document_key")
            if "::" in needle:
                key, needle = needle.split("::", 1)
            matches = [c["chunk_id"] for c in chunks if c["document_id"] == documents[key]["id"] and needle in " ".join(c["content"].split())]
            if len(matches) != 1:
                raise RuntimeError(f"Ambiguous/missing ground truth for {q['id']}: {needle!r}, {matches}")
            ids += matches
        q["gold_chunk_ids"] = sorted(set(ids))
    assert len(chunks) >= 25
    assert all(c["document_id"] != documents["private"]["id"] for c in chunks)
    assert all(c["document_version_id"] != documents["travel"]["historical_version_id"] for c in chunks)
    manifest = {"corpus_sha256": digest(corpus), "member_id": member["id"],
                "documents": documents, "chunks": chunks, "questions": questions}
    manifest["index_revision"] = digest(chunks)
    save("manifest.json", manifest)
    print("frozen_manifest_sha256", digest(manifest), "chunks", len(chunks), "questions", len(questions))
    admin.close()


def baseline():
    manifest = load("manifest.json")
    member = API("rag-experiment")
    rows = []
    # Warm up transport/SQL connection pool without using a benchmark question.
    member.search("прогревсоединения")
    for q in manifest["questions"]:
        elapsed = []
        for _ in range(3):
            start = time.perf_counter()
            sources = member.search(q["query"])
            elapsed.append((time.perf_counter() - start) * 1000)
        assert all(s["chunk_id"] in {c["chunk_id"] for c in manifest["chunks"]} for s in sources)
        rows.append({**q, "results": [(s["chunk_id"], s["score"]) for s in sources],
                     "latency_ms": float(np.median(elapsed)), "latency_samples_ms": elapsed})
    save("fts.json", rows)
    save("baseline-summary.json", {s: metrics([r for r in rows if r["split"] == s]) for s in ("development", "holdout")})
    member.close()
    print(json.dumps(load("baseline-summary.json"), indent=2))


class Resources:
    def __enter__(self):
        self.process = psutil.Process()
        self.initial_cpu = self.process.cpu_times()
        self.start = time.perf_counter()
        self.peak = self.process.memory_info().rss
        self.stop = threading.Event()
        def sample():
            while not self.stop.wait(0.02):
                self.peak = max(self.peak, self.process.memory_info().rss)
        self.thread = threading.Thread(target=sample, daemon=True)
        self.thread.start()
        return self

    def __exit__(self, *_):
        self.stop.set()
        self.thread.join()
        cpu = self.process.cpu_times()
        wall = time.perf_counter() - self.start
        seconds = cpu.user + cpu.system - self.initial_cpu.user - self.initial_cpu.system
        self.result = {"wall_seconds": wall, "cpu_seconds": seconds,
                       "cpu_percent_one_core": 100 * seconds / wall, "rss_peak_mib": self.peak / 2**20,
                       "rss_end_mib": self.process.memory_info().rss / 2**20}


def index():
    manifest = load("manifest.json")
    if visible_chunks(manifest["member_id"]) != manifest["chunks"]:
        raise RuntimeError("Visibility/version snapshot changed; recreate baseline")
    with Resources() as resources:
        start = time.perf_counter()
        model = E5(MODEL_PATH)
        load_ms = (time.perf_counter() - start) * 1000
        vectors, window_ids, spans = [], [], []
        for chunk in manifest["chunks"]:
            # Title is shared contextual data; content and source identity unchanged.
            values, windows = model.encode(chunk["document_title"] + "\n" + chunk["content"], "passage")
            vectors.extend(values)
            window_ids.extend([chunk["chunk_id"]] * len(values))
            spans.extend({"chunk_id": chunk["chunk_id"], **span} for span in windows)
        np.save(OUT / "passages.npy", np.asarray(vectors))
    metadata = {"model": MODEL, "model_revision": REVISION, "dimension": DIMENSION,
                "index_revision": manifest["index_revision"], "readiness": "ready",
                "model_sha256": file_digest(MODEL_PATH / "onnx/model.onnx"),
                "tokenizer_sha256": file_digest(MODEL_PATH / "tokenizer.json"),
                "prefixes": {"query": "query: ", "passage": "passage: "}, "normalized": True,
                "pooling": "attention-mask mean", "tokenization_revision": "e5-exact-prefix-v1",
                "chunk_score": "max cosine over all token windows",
                "window_overlap_tokens": 64, "threads": 2, "window_ids": window_ids,
                "spans": spans, "load_ms": load_ms, "resources": resources.result}
    save("index.json", metadata)
    print("indexed", len(window_ids), "windows", "long_chunks", len({s["chunk_id"] for s in spans if s["total_content_tokens"] > 508}), resources.result)


def compare():
    if (OUT / "frozen-parameters.json").exists():
        raise RuntimeError("Holdout already evaluated. Do not tune/re-run on the same holdout")
    manifest, metadata, baseline_rows = load("manifest.json"), load("index.json"), load("fts.json")
    if (metadata["model"], metadata["model_revision"], metadata["dimension"], metadata["index_revision"]) != (
            MODEL, REVISION, DIMENSION, manifest["index_revision"]):
        raise RuntimeError("Incompatible model/index")
    if metadata["tokenization_revision"] != "e5-exact-prefix-v1" or metadata["readiness"] != "ready":
        raise RuntimeError("Incompatible formatting or incomplete embedding index")
    if metadata["model_sha256"] != file_digest(MODEL_PATH / "onnx/model.onnx") or metadata["tokenizer_sha256"] != file_digest(MODEL_PATH / "tokenizer.json"):
        raise RuntimeError("Local model/tokenizer changed; explicit reindex required")
    if visible_chunks(manifest["member_id"]) != manifest["chunks"]:
        raise RuntimeError("ACL or active versions changed; index cannot be used")
    vectors = np.load(OUT / "passages.npy", allow_pickle=False)
    if vectors.shape != (len(metadata["window_ids"]), DIMENSION) or not np.isfinite(vectors).all():
        raise RuntimeError("Corrupt vector index")
    if set(metadata["window_ids"]) != {c["chunk_id"] for c in manifest["chunks"]}:
        raise RuntimeError("Incomplete index or unauthorized/unknown source IDs")
    assert np.allclose(np.linalg.norm(vectors, axis=1), 1, atol=1e-5)
    config = {"rrf_k": 60, "candidate_pool": 10, "semantic_cutoff": None,
              "tokenization_revision": "e5-exact-prefix-v1",
              "history": "current question only in all branches; contextual followups are a measured limitation",
              "score_threshold": "not calibrated: raw top-k compares retrieval recall and exposes negative evidence risk",
              "corpus_sha256": manifest["corpus_sha256"], "manifest_sha256": digest(manifest)}
    # Parameter-free comparison. Freeze before accessing holdout embeddings/results.
    save("frozen-parameters.json", config)
    rows, hybrid_rows = [], []
    by_query = {r["id"]: r for r in baseline_rows}
    with Resources() as resources:
        model = E5(MODEL_PATH)
        model.encode("прогрев вычислений", "query")
        for q in manifest["questions"]:
            elapsed = []
            for _ in range(3):
                start = time.perf_counter()
                encoded, _ = model.encode(q["query"], "query")
                candidates = cosine_candidates(encoded, vectors, metadata["window_ids"])[:10]
                elapsed.append((time.perf_counter() - start) * 1000)
            original = by_query[q["id"]]
            start = time.perf_counter()
            fused = rrf([original["results"], candidates], k=config["rrf_k"])
            fusion_ms = (time.perf_counter() - start) * 1000
            rows.append({**q, "results": candidates, "latency_ms": float(np.median(elapsed)),
                         "latency_samples_ms": elapsed})
            # Sequential local experiment; no claims of production parallel latency.
            hybrid_rows.append({**q, "results": fused,
                                "latency_ms": float(np.median(elapsed)) + original["latency_ms"] + fusion_ms,
                                "fusion_ms": fusion_ms})
    save("semantic.json", rows)
    save("hybrid.json", hybrid_rows)
    summary = {method: {s: metrics([r for r in results if r["split"] == s]) for s in ("development", "holdout")}
               for method, results in (("fts", baseline_rows), ("semantic", rows), ("hybrid", hybrid_rows))}
    summary["resources"] = resources.result
    save("comparison.json", summary)
    print(json.dumps(summary, indent=2))


def python_token():
    # Only the isolated Python fixture; do not serialize the service token into artifacts.
    container = json.loads(subprocess.check_output(["docker", "inspect", "hackathon-hybrid-test-ai"]))[0]
    return next(e.split("=", 1)[1] for e in container["Config"]["Env"] if e.startswith("AI_SERVICE_TOKEN="))


def answer_one(method, row, manifest, token, api=None, contexts=None):
    sources = {c["chunk_id"]: c for c in manifest["chunks"]}
    selected = [sources[sid] for sid, _ in row["results"][:8]]
    start = time.perf_counter()
    session_id = None
    if method == "fts":
        # Real public Chat API with original retrieval and unchanged Qwen generation.
        session_id = contexts[row["id"]]["session_id"] if row["id"] in contexts else api.request(
            "POST", "/api/v1/chat/sessions", json={"title": "Synthetic retrieval evaluation"})["id"]
        path = "/api/v1/chat/sessions/" + session_id
        turn = api.request("POST", path + "/messages",
                           json={"content": row["query"], "client_message_id": str(uuid4())})
        source_ids = [s["chunk_id"] for s in turn["sources"]]
        result = {"mode": turn["mode"], "answer": turn["message"], "used_source_ids": source_ids}
        api.request("DELETE", path)
    else:
        # Offline retrieval arms use exactly the existing private generation API,
        # never install a new retrieval path in Go/Python or alter its prompts.
        payload = {"query": row["query"], "history": contexts.get(row["id"], {}).get("history", []), "sources": [
            {key: s[key] for key in ("chunk_id", "document_title", "version_number", "section_number", "page_number", "content")}
            for s in selected]}
        with httpx.Client(trust_env=False, timeout=90) as client:
            response = client.post(PYTHON + "/internal/v1/rag/generate", json=payload,
                                   headers={"X-Service-Token": token})
            response.raise_for_status()
            body = response.json()
            allowed = {s["chunk_id"] for s in selected}
            if len(body["used_source_ids"]) != len(set(body["used_source_ids"])) or not set(body["used_source_ids"]) <= allowed:
                raise RuntimeError("Generation returned invalid citations")
            result = {"mode": body["status"], "answer": body["answer"] or body["clarification_question"] or "",
                      "used_source_ids": body["used_source_ids"]}
    result.update({"query_id": row["id"], "method": method, "split": row["split"], "query": row["query"],
                   "history": contexts.get(row["id"], {}).get("history", []),
                   "gold_chunk_ids": row["gold_chunk_ids"], "facts": row.get("facts", []),
                   "retrieved_chunk_ids": [s["chunk_id"] for s in selected],
                   "elapsed_ms": (time.perf_counter() - start) * 1000, "chat_session_id": session_id})
    gold, used = set(row["gold_chunk_ids"]), set(result["used_source_ids"])
    answered = result["mode"] in {"generated", "answered"}
    result["answered"] = answered
    result["citation_gold_precision"] = len(gold & used) / len(used) if used else None
    result["gold_fully_cited"] = bool(gold) and gold <= used
    # This is a protocol signal, not factual hallucination. An answered response
    # may explicitly say that a requested name/value is absent; review its text.
    result["negative_answered_status"] = not gold and answered
    # Mechanical screen only; report human-reviewed semantic correctness separately.
    result["fact_substring_screen"] = answered and all(f.casefold() in result["answer"].casefold() for f in row.get("facts", []))
    return result


def answers():
    manifest, token = load("manifest.json"), python_token()
    if visible_chunks(manifest["member_id"]) != manifest["chunks"]:
        raise RuntimeError("ACL/version snapshot changed; refusing to send stale sources to Qwen")
    rows = load("answers.json") if (OUT / "answers.json").exists() else []
    completed = {(r["method"], r["query_id"]) for r in rows}
    errors = []
    api = API("rag-experiment")
    contexts = load("context-history.json") if (OUT / "context-history.json").exists() else {}
    for question in manifest["questions"]:
        if not question.get("history") or question["id"] in contexts:
            continue
        session = api.request("POST", "/api/v1/chat/sessions", json={"title": "Synthetic follow-up evaluation"})
        path = "/api/v1/chat/sessions/" + session["id"]
        # Build actual safe history through Chat API. No fabricated assistant turn.
        for prior in question["history"]:
            if prior["role"] == "user":
                api.request("POST", path + "/messages", json={"content": prior["content"], "client_message_id": str(uuid4())})
        history = api.request("GET", path)["messages"]
        contexts[question["id"]] = {"session_id": session["id"], "history": [
            {"role": h["role"], "content": h["content"]} for h in history]}
    save("context-history.json", contexts)
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
        tasks = {pool.submit(answer_one, method, row, manifest, token, api, contexts): (method, row["id"])
                 for method in ("fts", "semantic", "hybrid") for row in load(method + ".json")
                 if (method, row["id"]) not in completed}
        for task in concurrent.futures.as_completed(tasks):
            method, qid = tasks[task]
            try:
                result = task.result()
                rows.append(result)
                print(method, qid, result["mode"], "cited", len(result["used_source_ids"]), flush=True)
            except Exception as exc:
                detail = {"method": method, "query_id": qid, "error_type": type(exc).__name__}
                if isinstance(exc, httpx.HTTPStatusError):
                    detail["http_status"] = exc.response.status_code
                    detail["body"] = exc.response.json()
                errors.append(detail)
                print(method, qid, type(exc).__name__, flush=True)
            save("answers.json", sorted(rows, key=lambda r: (r["method"], r["query_id"])))
            save("answer-errors.json", errors)
    api.close()
    summary = {}
    for method in ("fts", "semantic", "hybrid"):
        summary[method] = {}
        for split in ("development", "holdout"):
            group = [r for r in rows if r["method"] == method and r["split"] == split]
            if not group:
                summary[method][split] = {"completed": 0}
                continue
            pos, neg = [r for r in group if r["gold_chunk_ids"]], [r for r in group if not r["gold_chunk_ids"]]
            citations = [r["citation_gold_precision"] for r in group if r["citation_gold_precision"] is not None]
            summary[method][split] = {"completed": len(group), "negative_answered_status_rate": np.mean([r["negative_answered_status"] for r in neg]).item(),
                                     "fully_cited_gold_rate": np.mean([r["gold_fully_cited"] for r in pos]).item(),
                                     "citation_gold_precision": np.mean(citations).item() if citations else None,
                                     "fact_substring_screen_rate": np.mean([r["fact_substring_screen"] for r in pos]).item(),
                                     "generation_ms_p50": float(np.percentile([r["elapsed_ms"] for r in group], 50)),
                                     "generation_ms_p95": float(np.percentile([r["elapsed_ms"] for r in group], 95))}
    save("answer-summary.json", summary)
    if errors:
        raise RuntimeError(f"{len(errors)} failed generations, retained in answer-errors.json")


def verify():
    manifest = load("manifest.json")
    member, admin = API("rag-experiment"), API()
    # All snapshot downloads must be actual ACL-protected downloadable originals.
    for record in manifest["documents"].values():
        download = member.http.get(f'/api/v1/documents/{record["id"]}/download')
        status = download.status_code
        assert status == (404 if record["key"] == "private" else 200), (record["key"], status)
        if status == 200:
            assert hashlib.sha256(download.content).hexdigest() == record["sha256"]
    original = visible_chunks(manifest["member_id"])
    doc = manifest["documents"]["security"]
    admin.request("DELETE", f'/api/v1/documents/{doc["id"]}/access/{manifest["member_id"]}')
    after_revoke = visible_chunks(manifest["member_id"])
    assert not any(c["document_id"] == doc["id"] for c in after_revoke)
    assert not member.search("компрометация")
    admin.request("POST", f'/api/v1/documents/{doc["id"]}/access', json={"user_id": manifest["member_id"]})
    assert visible_chunks(manifest["member_id"]) == original
    # Existing FTS remains operational even if optional Python service is stopped.
    before = member.search("компрометация")
    subprocess.run(["docker", "stop", "hackathon-hybrid-test-ai"], check=True, capture_output=True)
    try:
        assert member.search("компрометация") == before
        session = member.request("POST", "/api/v1/chat/sessions", json={"title": "Synthetic outage check"})
        path = "/api/v1/chat/sessions/" + session["id"]
        payload = {"content": "компрометация", "client_message_id": str(uuid4())}
        turn = member.request("POST", path + "/messages", json=payload)
        assert turn["mode"] == "search" and len(turn["sources"]) == 1
        assert member.request("POST", path + "/messages", json=payload) == turn
        member.request("DELETE", path)
    finally:
        subprocess.run(["docker", "start", "hackathon-hybrid-test-ai"], check=True, capture_output=True)
    # Restart only disposable fixtures, retaining their container storage.
    subprocess.run(["docker", "restart", PG], check=True, capture_output=True)
    subprocess.run(["docker", "restart", "hackathon-hybrid-test-api"], check=True, capture_output=True)
    for _ in range(60):
        try:
            if visible_chunks(manifest["member_id"]) == original:
                break
        except (subprocess.CalledProcessError, json.JSONDecodeError):
            pass
        time.sleep(0.5)
    assert visible_chunks(manifest["member_id"]) == original
    for _ in range(60):
        try:
            if member.http.get("/healthz").status_code == 200:
                break
        except httpx.RequestError:
            pass
        time.sleep(0.5)
    assert member.search("компрометация") == before
    save("security-checks.json", {"downloads": True, "sql_acl_revoke": True, "fts_acl_revoke": True,
                                  "historical_and_private_excluded": True, "fts_python_outage": True,
                                  "download_sha256_matches": True, "chat_outage_search_fallback": True,
                                  "chat_idempotency": True, "restart_chunk_durability": True,
                                  "restart_sessions_and_search": True, "semantic_is_offline_snapshot_only": True})
    admin.close()
    member.close()
    print("security/download/outage/restart checks passed")


def probe():
    """A long, synthetic safety probe, separate from the frozen retrieval corpus."""
    model = E5(MODEL_PATH)
    passage = ("Проверка журнала технического обслуживания. " * 250) + " При аварии аварийный ключ находится в опечатанном шкафу номер 73."
    vectors, spans = model.encode(passage, "passage")
    assert max(s["input_tokens"] for s in spans) <= 512
    assert spans[-1]["end_token"] == spans[-1]["total_content_tokens"]
    try:
        model.encode(passage, "query")
    except ValueError:
        pass
    else:
        raise AssertionError("Overlong query truncated silently")
    query, _ = model.encode("Где находится аварийный ключ?", "query")
    scores = (vectors @ query[0]).tolist()
    save("long-context-probe.json", {"total_tokens": spans[-1]["total_content_tokens"], "spans": spans,
                                     "cosine_per_window": scores, "best_window": scores.index(max(scores)),
                                     "normalized": bool(np.allclose(np.linalg.norm(vectors, axis=1), 1, atol=1e-5)),
                                     "overlong_query_rejected": True})
    print("long-context probe passed; windows", len(spans), "best window", scores.index(max(scores)))


def environment():
    images = {}
    for name in ("hackathon-hybrid-test-api", "hackathon-hybrid-test-ai", PG):
        container = json.loads(subprocess.check_output(["docker", "inspect", name]))[0]
        # Never persist Config.Env, inference URLs, tokens, mounts, or credentials.
        images[name] = {"image_id": container["Image"], "memory_limit_bytes": container["HostConfig"]["Memory"],
                        "nano_cpus": container["HostConfig"]["NanoCpus"]}
    packages = {p: importlib.metadata.version(p) for p in ("numpy", "onnxruntime", "tokenizers", "httpx", "psutil", "python-docx", "PyMuPDF")}
    save("environment.json", {"platform": platform.platform(), "python": platform.python_version(),
                              "logical_cpu_count": os.cpu_count(), "host_ram_mib": psutil.virtual_memory().total / 2**20,
                              "packages": packages, "containers": images,
                              "git_head": subprocess.check_output(["git", "rev-parse", "HEAD"], text=True).strip()})


def diagnostics():
    manifest, semantic, fts = load("manifest.json"), load("semantic.json"), load("fts.json")
    chunks = {c["chunk_id"]: c for c in manifest["chunks"]}
    per_document = []
    for key, document in manifest["documents"].items():
        group = [r for r in semantic if r["gold_chunk_ids"] and any(
            chunks[sid]["document_id"] == document["id"] for sid in r["gold_chunk_ids"])]
        if not group:
            continue
        ids = {r["id"] for r in group}
        per_document.append({"key": key, "query_count": len(group),
                             "fts": metrics([r for r in fts if r["id"] in ids]), "semantic": metrics(group)})
    save("per-document.json", per_document)
    dev = [r for r in semantic if r["split"] == "development"]
    pos = [r["results"][0][1] for r in dev if r["gold_chunk_ids"]]
    neg = [r["results"][0][1] for r in dev if not r["gold_chunk_ids"]]
    save("score-overlap.json", {"development_positive_top1_range": [min(pos), max(pos)],
                                "development_negative_top1_range": [min(neg), max(neg)], "no_threshold_fitted": True})
    model = E5(MODEL_PATH)
    vectors, metadata = np.load(OUT / "passages.npy", allow_pickle=False), load("index.json")
    misses = []
    for row in semantic:
        missing = set(row["gold_chunk_ids"]) - {sid for sid, _ in row["results"]}
        if not missing:
            continue
        query, _ = model.encode(row["query"], "query")
        candidates = cosine_candidates(query, vectors, metadata["window_ids"])
        for sid in sorted(missing):
            rank = next(i + 1 for i, item in enumerate(candidates) if item[0] == sid)
            misses.append({"query_id": row["id"], "gold_id": sid, "rank_in_full_corpus": rank,
                           "cosine": candidates[rank - 1][1], "no_parameter_changes": True})
    save("miss-analysis.json", misses)


def export():
    """Copy an explicit allowlist of synthetic results, never environments/models."""
    destination = HERE / "results/2026-10-24"
    destination.mkdir(parents=True, exist_ok=True)
    names = ["manifest.json", "fts.json", "semantic.json", "hybrid.json", "comparison.json",
             "frozen-parameters.json", "index.json", "answers.json", "answer-errors.json",
             "review.json", "review-summary.json", "per-document.json", "score-overlap.json",
             "miss-analysis.json", "long-context-probe.json", "security-checks.json",
             "environment.json", "container-stats.json", "context-history.json"]
    # Check both secret and endpoint absence without printing their values.
    container = json.loads(subprocess.check_output(["docker", "inspect", "hackathon-hybrid-test-ai"]))[0]
    private_values = [e.split("=", 1)[1] for e in container["Config"]["Env"] if
                      e.startswith(("AI_SERVICE_TOKEN=", "LLM_API_KEY=", "LLM_BASE_URL=")) and e.split("=", 1)[1]]
    for name in names:
        raw = (OUT / name).read_text(encoding="utf-8")
        if any(value in raw for value in private_values):
            raise RuntimeError("Refusing export: private configuration found")
        shutil.copyfile(OUT / name, destination / name)
    print("exported", len(names), "synthetic JSON artifacts; private configuration scan passed")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("stage", choices=["prepare", "baseline", "index", "compare", "answers", "verify", "probe", "environment", "diagnostics", "export"])
    globals()[parser.parse_args().stage]()
