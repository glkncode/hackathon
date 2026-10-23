import io
import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from main import app
from app.templates.engine import MAX_TEMPLATE_BYTES
from app.templates.routes import DOCX_MIME

TOKEN = "template-test-token"


def make_template(text="Кому: {{ recipient }}"):
    document = Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def edit_zip(content, replacements):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            if entry.filename not in replacements:
                target.writestr(entry.filename, source.read(entry))
        for name, data in replacements.items():
            target.writestr(name, data)
    return output.getvalue()


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("AI_SERVICE_TOKEN", TOKEN)
    monkeypatch.delenv("LLM_BASE_URL", raising=False)
    with TestClient(app) as client:
        yield client


def send(client, action="inspect", content=None, values=None, token=TOKEN, filename="memo.docx"):
    data = {} if values is None else {"values": values if isinstance(values, str) else json.dumps(values)}
    return client.post("/internal/v1/templates/" + action, headers={"X-Service-Token": token},
                       data=data, files={"template": (filename, make_template() if content is None else content, DOCX_MIME)})


def test_inspect_missing_and_filled(client):
    response = send(client)
    assert response.status_code == 200
    assert response.json() == {
        "fields": [{"name": "recipient", "type": "string", "required": True}],
        "required_fields": ["recipient"], "missing_fields": ["recipient"], "can_render": False,
    }
    response = send(client, values={"recipient": "Иван"})
    assert response.json()["can_render"] is True
    assert response.json()["missing_fields"] == []


def test_render_valid_docx_xml_escaping(client):
    value = 'Иван <Петров> & "Партнёры"'
    response = send(client, "render", values={"recipient": value})
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_MIME
    assert response.headers["cache-control"] == "no-store"
    assert "attachment" in response.headers["content-disposition"]
    assert Document(io.BytesIO(response.content)).paragraphs[0].text == "Кому: " + value


def test_inspect_split_runs_header_footer_table(client):
    document = Document()
    paragraph = document.add_paragraph()
    for fragment in ["{{ rec", "ipient", " }}"]:
        paragraph.add_run(fragment)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "{{ body }}"
    document.sections[0].header.paragraphs[0].text = "{{ header }}"
    document.sections[0].footer.paragraphs[0].text = "{{ footer }}"
    content = io.BytesIO()
    document.save(content)
    assert send(client, content=content.getvalue()).json()["required_fields"] == ["body", "footer", "header", "recipient"]
    result = send(client, "render", content=content.getvalue(), values={
        "body": "Таблица", "footer": "Подвал", "header": "Шапка", "recipient": "Адресат",
    })
    assert result.status_code == 200
    parsed = Document(io.BytesIO(result.content))
    assert parsed.paragraphs[0].text == "Адресат"
    assert parsed.tables[0].cell(0, 0).text == "Таблица"
    assert parsed.sections[0].header.paragraphs[0].text == "Шапка"
    assert parsed.sections[0].footer.paragraphs[0].text == "Подвал"


@pytest.mark.parametrize("values", [{}, {"recipient": ""}, {"recipient": "  "}])
def test_render_missing_fields(client, values):
    response = send(client, "render", values=values)
    assert response.status_code == 422
    assert response.json() == {"error": {"code": "missing_fields", "missing_fields": ["recipient"]}}


@pytest.mark.parametrize("action", ["inspect", "render"])
@pytest.mark.parametrize("token", ["", "wrong"])
def test_authentication(client, action, token):
    assert send(client, action, token=token).status_code == 401


def test_auth_before_body(client):
    assert client.post("/internal/v1/templates/inspect", content=b"broken").status_code == 401


@pytest.mark.parametrize("content,filename,status", [
    (b"", "memo.docx", 422), (b"not a zip", "memo.docx", 422),
    (b"hello", "memo.txt", 415), (b"x" * (MAX_TEMPLATE_BYTES + 1), "memo.docx", 413),
], ids=["empty", "bad-zip", "wrong-extension", "too-large"])
def test_bad_files(client, content, filename, status):
    assert send(client, content=content, filename=filename).status_code == status


@pytest.mark.parametrize("raw", [
    "not json", "[]", '{"recipient":123}', '{"recipient":null}',
    '{"recipient":"a","recipient":"b"}', '{"recipient":"\\u0000"}',
    '{"recipient":"\\ud800"}', '{"bad.key":"text"}',
])
def test_bad_values(client, raw):
    response = send(client, values=raw)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_values"


def test_unknown_field(client):
    response = send(client, values={"typo": "someone"})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "unknown_fields"


@pytest.mark.parametrize("syntax", [
    "{{ recipient.__class__ }}", "{{ recipient|upper }}", "{{ 7*7 }}",
    "{% for i in range(99999999) %}x{% endfor %}", "{{ cycler.__init__.__globals__ }}",
    "{% include '/etc/passwd' %}", "{# comment #}", "{{r recipient }}", "{{ invalid",
])
def test_disallow_executable_templates(client, syntax):
    assert send(client, content=make_template(syntax)).status_code == 422


def test_values_not_reinterpreted_as_jinja(client):
    value = "{{ cycler.__init__.__globals__ }}"
    response = send(client, "render", values={"recipient": value})
    assert response.status_code == 200
    assert value in Document(io.BytesIO(response.content)).paragraphs[0].text


def test_external_relationship(client):
    rel = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId999" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid" TargetMode="External"/></Relationships>'
    response = send(client, content=edit_zip(make_template(), {"word/_rels/document.xml.rels": rel}))
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "unsafe_template"


@pytest.mark.parametrize("name,data", [
    ("../outside.xml", b"x"), ("word/vbaProject.bin", b"x"),
    ("word/embeddings/oleObject.bin", b"x"),
    ("custom.xml", b'<!DOCTYPE x [<!ENTITY a SYSTEM "file:///etc/passwd">]><x>&a;</x>'),
    ("custom.xml", b"<x>{{ hidden }}</x>"),
])
def test_unsafe_zip_parts(client, name, data):
    assert send(client, content=edit_zip(make_template(), {name: data})).status_code == 422


def test_zip_expanded_size_limit(client):
    content = edit_zip(make_template(), {"large.txt": b"x" * (32 * 1024 * 1024)})
    assert send(client, content=content).status_code == 413


def test_parallel_requests_are_isolated(client):
    def render(index):
        response = send(client, "render", values={"recipient": str(index)})
        assert response.status_code == 200
        return Document(io.BytesIO(response.content)).paragraphs[0].text
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(render, range(8)))
    assert results == ["Кому: " + str(i) for i in range(8)]


def test_demo(client):
    root = Path(__file__).resolve().parents[1] / "examples"
    content = (root / "memo-template.docx").read_bytes()
    values = json.loads((root / "memo-values.json").read_text(encoding="utf-8"))
    response = send(client, content=content, values=values)
    assert response.status_code == 200
    assert response.json()["can_render"] is True
    assert len(response.json()["fields"]) == 10
    response = send(client, "render", content=content, values=values)
    assert response.status_code == 200
    paragraphs = "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert values["subject"] in paragraphs
    assert values["request_text"] in paragraphs
    assert "{{" not in paragraphs


def test_values_limit(client):
    raw = json.dumps({f"field_{i}": "Я" * 6000 for i in range(6)}, ensure_ascii=False)
    response = send(client, values=raw)
    assert response.status_code == 413


def test_whole_request_limit(client):
    response = client.post(
        "/internal/v1/templates/inspect",
        headers={"X-Service-Token": TOKEN, "Content-Type": "multipart/form-data; boundary=limit"},
        content=iter([
            b'--limit\r\nContent-Disposition: form-data; name="template"; filename="a.docx"\r\n\r\n',
            *([b"x" * (1024 * 1024)] * 7), b"\r\n--limit--\r\n",
        ]),
    )
    assert response.status_code == 413


def test_constant_template(client):
    content = make_template("Постоянный текст")
    assert send(client, content=content).json()["fields"] == []
    assert send(client, "render", content=content, values={}).status_code == 200


def test_render_requires_values_form(client):
    response = send(client, "render")
    assert response.status_code == 422
    assert response.json() == {"error": {"code": "invalid_template_request"}}
