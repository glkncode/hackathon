import io
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from app.templates.engine import TemplateError, process_template


def schema(paragraph=0, match="________", kind="text", required=True):
    return json.dumps([dict(key="value", label="Значение", type=kind, required=required,
                           question="Укажите значение", max_length=100,
                           bindings=[dict(part="word/document.xml", paragraph=paragraph, match=match)])])


def fixture_doc():
    d = Document()
    p = d.add_paragraph()
    p.add_run("До ").bold = True
    p.add_run("____").italic = True
    p.add_run("____ после")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Участники: ________"
    d.add_paragraph("НЕ УДАЛЯТЬ")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def test_mapped_split_runs_styles_parts_original():
    original = fixture_doc()
    result = process_template(original, json.dumps({"value": "Текст & <имя>"}), True, schema())
    d = Document(io.BytesIO(result))
    assert d.paragraphs[0].text == "До Текст & <имя> после"
    assert d.paragraphs[0].runs[0].bold
    assert d.paragraphs[0].runs[1].italic
    assert d.paragraphs[1].text == "НЕ УДАЛЯТЬ"
    assert Document(io.BytesIO(original)).paragraphs[0].text == "До ________ после"
    with ZipFile(io.BytesIO(original)) as a, ZipFile(io.BytesIO(result)) as b:
        for name in a.namelist():
            if name != "word/document.xml":
                assert a.read(name) == b.read(name)


def test_mapped_repeated_people_table():
    result = process_template(fixture_doc(), json.dumps({"value": ["Иванов", "Петров"]}), True,
                              schema(1, kind="person_list"))
    assert Document(io.BytesIO(result)).tables[0].cell(0, 0).text == "Участники: Иванов\nПетров"


@pytest.mark.parametrize("value,kind", [(" ", "text"), ([], "person_list"),
                                         ("2026-02-30", "date"), ("x" * 101, "text"), (42, "text")])
def test_mapped_invalid(value, kind):
    with pytest.raises(TemplateError):
        process_template(fixture_doc(), json.dumps({"value": value}), True, schema(kind=kind))


def test_optional_and_binding_guard():
    result = process_template(fixture_doc(), "{}", True, schema(required=False))
    assert Document(io.BytesIO(result)).paragraphs[0].text == "До  после"
    with pytest.raises(TemplateError):
        process_template(fixture_doc(), "{}", False, schema(2, "НЕ УДАЛЯТЬ"))


def test_real_working_group():
    root = Path(__file__).resolve().parents[2]
    files = list((root / "data/demo").glob("*рабоч*группы.docx"))
    if not files:
        pytest.skip("real source DOCX is in repository data/demo")
    original = files[0].read_bytes()
    fields = json.loads((root / "docs/template-profiles/working-group.json").read_text(encoding="utf-8-sig"))
    values = {f["key"]: "Тестовое значение " + f["key"] for f in fields}
    values.update(members=["Сидоров С.С., специалист ИТ", "Кузнецова Е.В., архив"],
                  deadline="2026-12-15", approval_date="2026-10-24",
                  group_subject="Внедрение электронного архива")
    result = process_template(original, json.dumps(values), True, json.dumps(fields))
    d = Document(io.BytesIO(result))
    assert d.tables[2].cell(3, 1).text == "\n".join(values["members"])
    assert "15.12.2026" in d.paragraphs[9].text
    assert values["group_subject"] in d.paragraphs[5].text
    with ZipFile(io.BytesIO(original)) as a, ZipFile(io.BytesIO(result)) as b:
        for name in a.namelist():
            if name != "word/document.xml":
                assert a.read(name) == b.read(name)
    assert "ВИЗУАЛИЗАЦИЯ ЭП (НЕ УДАЛЯТЬ)" in d.tables[3].cell(1, 0).text
