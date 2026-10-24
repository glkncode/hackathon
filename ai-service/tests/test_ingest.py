from io import BytesIO
from uuid import uuid4

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from main import app
from app.chunking.text import make_chunks
from app.core.config import MAX_FILE_BYTES
from app.ingestion.errors import IngestionError

TOKEN = "test-private-token"
DOCUMENT_ID = str(uuid4())
VERSION_ID = str(uuid4())


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("AI_SERVICE_TOKEN", TOKEN)
    with TestClient(app) as client:
        yield client


def upload(client, name, content, mime="text/plain", token=TOKEN, **fields):
    return client.post(
        "/internal/v1/ingest",
        headers={"X-Service-Token": token},
        data={"document_id": DOCUMENT_ID, "document_version_id": VERSION_ID,
              "mime_type": mime, **fields},
        files={"file": (name, content)},
    )


def test_txt_contract_and_determinism(client):
    content = "1.2 Общие положения\nТекст документа\n" + "Я" * 9000
    first = upload(client, "sample.txt", content.encode())
    assert first.status_code == 200
    body = first.json()
    assert set(body) == {"document_version_id", "chunks", "warnings"}
    assert body["document_version_id"] == VERSION_ID
    assert len(body["chunks"]) >= 3
    for index, chunk in enumerate(body["chunks"]):
        assert chunk["chunk_index"] == index
        assert chunk["page_number"] is None
        assert chunk["section_number"] == "1.2"
        assert chunk["section_title"] == "Общие положения"
        assert 0 < len(chunk["content"].encode()) <= 64 * 1024
    assert body == upload(client, "sample.txt", content.encode()).json()


def test_pdf_preserves_pages(client):
    with pymupdf.open() as pdf:
        for text in ["Page one", "Page two"]:
            pdf.new_page().insert_text((72, 72), text)
        content = pdf.tobytes()
    response = upload(client, "sample.pdf", content, "application/pdf")
    assert response.status_code == 200
    chunks = response.json()["chunks"]
    assert [chunk["page_number"] for chunk in chunks] == [1, 2]
    assert [chunk["content"] for chunk in chunks] == ["Page one", "Page two"]


def test_docx_with_table(client):
    doc = Document()
    doc.add_paragraph("Before")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table"
    doc.add_paragraph("After")
    buf = BytesIO()
    doc.save(buf)
    response = upload(client, "sample.docx", buf.getvalue(),
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.status_code == 200
    chunks = response.json()["chunks"]
    assert [c["content"] for c in chunks] == ["Before", "Table", "After"]
    assert all(c["page_number"] is None for c in chunks)


@pytest.mark.parametrize("token", ["", "wrong-token"])
def test_invalid_token(client, token):
    assert upload(client, "sample.txt", b"hello", token=token).status_code == 401


def test_missing_token_before_body_validation(client):
    assert client.post("/internal/v1/ingest", content=b"invalid").status_code == 401


@pytest.mark.parametrize("name,content,mime,status", [
    ("empty.txt", b"", "text/plain", 422),
    ("blank.txt", b" \n\t", "text/plain", 422),
    ("bad.pdf", b"broken", "application/pdf", 422),
    ("bad.docx", b"broken", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", 422),
    ("bad.txt", b"\xff\xfe", "text/plain", 422),
    ("bad.exe", b"hello", "text/plain", 415),
    ("bad.txt", b"hello", "application/pdf", 415),
])
def test_invalid_files(client, name, content, mime, status):
    assert upload(client, name, content, mime).status_code == status


def test_invalid_uuid(client):
    assert upload(client, "a.txt", b"hello", document_version_id="bad").status_code == 422


def test_file_limit(client):
    assert upload(client, "large.txt", b"x" * (MAX_FILE_BYTES + 1)).status_code == 413


def test_chunked_upload_without_content_length(client):
    boundary = "test-boundary"
    body = b""
    for name, value in [
        ("document_id", DOCUMENT_ID),
        ("document_version_id", VERSION_ID),
        ("mime_type", "text/plain"),
    ]:
        body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"{name}\"\r\n\r\n{value}\r\n").encode()
    body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"a.txt\"\r\nContent-Type: text/plain\r\n\r\nhello\r\n--{boundary}--\r\n").encode()
    response = client.post(
        "/internal/v1/ingest",
        headers={"X-Service-Token": TOKEN,
                 "Content-Type": f"multipart/form-data; boundary={boundary}"},
        content=(body[i:i + 19] for i in range(0, len(body), 19)),
    )
    assert response.status_code == 200
    assert response.json()["chunks"][0]["content"] == "hello"


def test_nul_removed():
    assert make_chunks([(None, "hello\x00world")])[0]["content"] == "helloworld"


def test_streamed_body_limit(client):
    response = client.post(
        "/internal/v1/ingest",
        headers={"X-Service-Token": TOKEN,
                 "Content-Type": "multipart/form-data; boundary=limit"},
        content=iter([
            b'--limit\r\nContent-Disposition: form-data; name="file"; filename="a.txt"\r\n\r\n',
            *([b"x" * (1024 * 1024)] * 22),
            b"\r\n--limit--\r\n",
        ]),
    )
    assert response.status_code == 413


def test_output_limits():
    with pytest.raises(IngestionError) as error:
        make_chunks([(None, "x" * (8 * 1024 * 1024 + 1))])
    assert error.value.status_code == 422
    with pytest.raises(IngestionError) as error:
        make_chunks([(None, "hello")] * 2001)
    assert error.value.status_code == 422


def test_missing_config_fails_closed(monkeypatch):
    monkeypatch.delenv("AI_SERVICE_TOKEN", raising=False)
    with pytest.raises(RuntimeError, match="AI_SERVICE_TOKEN"):
        with TestClient(app):
            pass
