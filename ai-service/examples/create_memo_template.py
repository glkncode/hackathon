"""Rebuild the synthetic memo template; run from any working directory."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


def build(destination: Path):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.15
    title = document.styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(16)
    # Remove inherited theme fonts and decorative title rules from the base file.
    for style in document.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for style in [normal, title]:
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if "Theme" in key:
                del fonts.attrib[key]
        for script in ["ascii", "hAnsi", "eastAsia", "cs"]:
            fonts.set(qn("w:" + script), "Times New Roman")
    document.add_paragraph("{{ organization_name }}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in ["Кому: {{ recipient_position }}", "{{ recipient_name }}",
                 "От: {{ author_position }}", "{{ author_name }}"]:
        paragraph = document.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_after = Pt(3)
    document.add_paragraph("Служебная записка", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Дата: {{ document_date }}")
    paragraph = document.add_paragraph()
    paragraph.add_run("Тема: ").bold = True
    paragraph.add_run("{{ subject }}")
    document.add_paragraph("{{ request_text }}")
    document.add_paragraph("Обоснование: {{ justification }}")
    document.add_paragraph("Срок исполнения: {{ due_date }}")
    signature = document.add_paragraph("{{ author_name }}    __________________")
    signature.paragraph_format.space_before = Pt(20)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = "Служебная записка"
    document.core_properties.subject = ""
    document.core_properties.comments = ""
    for paragraph in document.paragraphs:
        for border in list(paragraph._p.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


if __name__ == "__main__":
    build(Path(__file__).resolve().parent / "memo-template.docx")
